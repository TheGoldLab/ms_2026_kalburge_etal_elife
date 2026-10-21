"""Table 1 of the paper: task settings of every cohort and block, as a formatted .xlsx.

One row per block: the cohort (with the code's task name), the participants
who started and who finished every block, the step mean and SD, the reward
rule (coins per correct choice, coins and steps lost per error), steps per
block, step duration, and the bonus rule.

Sources: the deployed settings were reconstructed from the raw data's tc_*
and ssc_step_* columns (task/OL-MX-PD/readme.md for cohorts 1, 2 and 4; the
MXSC files' tc_* columns for cohort 3; every file in a cohort agrees). The
reward rules are checked against pigeon.simulation.BLOCK_DEFAULTS /
BLOCK_DEFAULTS_PD, which the analyses use. Step durations are UPDATE_INTERVAL
(200 ms) and PREDEFINED_UPDATE_INTERVAL (100 ms, PD) in the task code. Every
trial costs 1 step to start in every cohort.

Writes ../Figures/Table01.xlsx (next to the figures). Run from the project root:

    uv run python scripts/make_table1.py
"""

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from pigeon.data import get_data_table
from pigeon.path_defaults import FIGURES_DIR
from pigeon.simulation import BLOCK_DEFAULTS, BLOCK_DEFAULTS_PD

OUT_FILE = FIGURES_DIR / "Table01.xlsx"

# (cohort, task code, description, step duration in ms, rows); each row is
# (block, step mean, step SD, coins gained, coins lost, steps lost, steps per block, bonus).
_PEN = [(1, 0, 0), (1, 4, 0), (1, 0, 30)]  # no penalty, coin penalty, step penalty
COHORTS = [
    (
        1,
        "OL",
        "Blocked SNR",
        200,
        [
            (b, m, 0.15, *_PEN[(b - 1) % 3], 600, f"{t} coins")
            for b, m, t in [
                (1, "±0.05", 125),
                (2, "±0.05", 23),
                (3, "±0.05", 30),
                (4, "±0.15", 170),
                (5, "±0.15", 115),
                (6, "±0.15", 105),
            ]
        ],
    ),
    (
        2,
        "MX",
        "Mixed SNR",
        200,
        [
            (b, m, 0.15, *_PEN[(b - 1) % 3], 600, f"{t} coins")
            for b, m, t in [
                (1, "±0.05 or ±0.15 ᵃ", 150),
                (2, "±0.05 or ±0.15 ᵃ", 45),
                (3, "±0.05 or ±0.15 ᵃ", 50),
                (4, "±0.10", 150),
                (5, "±0.10", 70),
                (6, "±0.10", 65),
            ]
        ],
    ),
    (
        3,
        "MXSC",
        "Change point",
        200,
        [
            (1, "±0.01", 0.05, 1, 1, 0, 1000, "100 coins"),
            (2, "±0.03", 0.05, 1, 1, 0, 1000, "100 coins"),
            (3, "±0.01 → ±0.03 ᵇ", 0.05, 1, 1, 0, 1000, "100 coins"),
            (4, "±0.03 → ±0.01 ᵇ", 0.05, 1, 1, 0, 1000, "100 coins"),
        ],
    ),
    (
        4,
        "PD",
        "Predefined bound ᶜ",
        100,
        [
            (b, m, 0.15, *_PEN[(b - 1) % 3], 800, f"bound {r}")
            for b, m, r in [
                (1, "±0.05", "0–0.0025"),
                (2, "±0.05", "0.40–0.60"),
                (3, "±0.05", "0.265–0.375"),
                (4, "±0.15", "0–0.01"),
                (5, "±0.15", "0.13–0.20"),
                (6, "±0.15", "0.19–0.255"),
            ]
        ],
    ),
]
FOOTNOTES = [
    (
        "Every trial cost 1 step to start. A bonus was paid for each block in which the participant reached the "
        "target number of coins (cohort 4: set a bound in the given range)."
    ),
    "ᵃ Each trial's step mean was drawn at random from the two values.",
    "ᵇ The step mean changed within each trial, at a change point that varied from trial to trial.",
    "ᶜ Participants set the bound before each trial; the trial ended when the pigeon reached it.",
]
HEADER = [
    "Cohort",
    "Task (code)",
    "n started / finished",
    "Block",
    "Step mean",
    "Step SD",
    "Coins per correct",
    "Coins lost per error",
    "Steps lost per error",
    "Steps per block",
    "Step duration (ms)",
    "Bonus target",
]


def _check_rules():
    """The reward rules above match the ones the analyses use."""
    for cohort, task, _, _, rows in COHORTS:
        if task == "MXSC":
            continue
        defaults = BLOCK_DEFAULTS_PD if task == "PD" else BLOCK_DEFAULTS
        for block, _, _, gained, lost, steps_lost, steps, _ in rows:
            d = defaults[block]
            assert (gained, lost, steps_lost, steps) == (
                d["coins_gained"],
                d["coins_lost"],
                d["steps_lost"],
                d["steps_per_block"],
            ), (cohort, block)


def get_cohort_counts(task, last_block):
    """Participants with any data, and those with trials in the last block."""
    d = get_data_table(task_type=task, combine_snr=False)
    n_started = d["subject_index"].nunique()
    n_finished = d.loc[d["block_index"] == last_block, "subject_index"].nunique()
    return n_started, n_finished


def write_table(counts, out_file=OUT_FILE):
    """Write the formatted table to out_file."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Table 1"
    thin = Side(style="thin")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.append(["Table 1: Task settings for each cohort and block"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append(HEADER)
    for cell in ws[2]:
        cell.font = Font(bold=True)
        cell.alignment = center
        cell.border = Border(top=thin, bottom=thin)
    for cohort, task, desc, step_ms, rows in COHORTS:
        first = ws.max_row + 1
        for block, mean, sd, gained, lost, steps_lost, steps, bonus in rows:
            ws.append(
                [
                    cohort,
                    f"{desc} ({task})",
                    "{} / {}".format(*counts[task]),
                    block,
                    mean,
                    sd,
                    gained,
                    lost,
                    steps_lost,
                    steps,
                    step_ms,
                    bonus,
                ]
            )
        last = ws.max_row
        for col in (1, 2, 3, 11):  # one value per cohort
            ws.merge_cells(start_row=first, end_row=last, start_column=col, end_column=col)
        for row in ws.iter_rows(min_row=first, max_row=last):
            for cell in row:
                cell.alignment = center
        for cell in ws[last]:
            cell.border = Border(bottom=thin)
    ws.append([])
    for note in FOOTNOTES:
        ws.append([note])
        ws.merge_cells(start_row=ws.max_row, end_row=ws.max_row, start_column=1, end_column=len(HEADER))
        ws.cell(ws.max_row, 1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[ws.max_row].height = 30
    widths = [8, 22, 12, 7, 18, 8, 10, 10, 10, 10, 10, 16]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[2].height = 32
    out_file.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_file)


def main():
    _check_rules()
    counts = {task: get_cohort_counts(task, max(r[0] for r in rows)) for _, task, _, _, rows in COHORTS}
    for task, (n_started, n_finished) in counts.items():
        print(f"{task}: {n_started} participants with data, {n_finished} with the last block")
    write_table(counts)
    print(f"Saved {OUT_FILE}")
    print(pd.DataFrame([r for *_, rows in COHORTS for r in rows]).to_string(header=False, index=False))


if __name__ == "__main__":
    main()
